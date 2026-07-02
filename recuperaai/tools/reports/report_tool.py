from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any

from recuperaai.database.models import Client, Document, Session
from recuperaai.tools.base import ToolBase


SEVERITY_LABELS = {
    'critical': 'Crítico',
    'warning': 'Atenção',
    'info': 'Informativo',
}
SEVERITY_ORDER = {'critical': 1, 'warning': 2, 'info': 3}
SEVERITY_WEIGHTS = {'critical': 5, 'warning': 2, 'info': 1}


def _now_label() -> str:
    return datetime.now().strftime('%d/%m/%Y %H:%M')


def _timestamp() -> str:
    return datetime.now().strftime('%Y%m%d_%H%M%S')


def _safe_filename(value: str) -> str:
    clean = re.sub(r'[^\w\-.]+', '_', value.strip(), flags=re.UNICODE).strip('_')
    return clean[:80] or 'cliente'


def _money(value: object) -> str:
    if value in (None, ''):
        return 'R$ 0,00'
    try:
        return f'R$ {float(value):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    except Exception:
        return str(value)


def _plain(value: object) -> str:
    return '' if value is None else str(value)


def _doc_to_dict(document: Document) -> dict[str, Any]:
    return asdict(document)


def _client_to_dict(client: Client) -> dict[str, Any]:
    return asdict(client)


class ReportTool(ToolBase):
    name = 'reports'
    label = 'Relatórios'

    def build_client_report_data(self, session: Session, client_id: int) -> dict[str, Any]:
        """Monta um dossiê estruturado para PDF/Excel sem depender da UI.

        O relatório profissional usa uma única fonte consolidada de dados para
        evitar divergência entre PDF, Excel e tela. Ele reúne cliente, documentos,
        apontamentos abertos, resumo executivo, ranking e evidências por nota.
        """
        self.engine.auth.require(session, 'reports_export')
        client = self.engine.clients.get(client_id)
        if not client:
            raise ValueError('Cliente não encontrado.')

        docs = self.engine.documents.list_for_client(client_id)
        invoices = [doc for doc in docs if doc.category == 'invoice']
        client_files = [doc for doc in docs if doc.category != 'invoice']
        findings = self.engine.findings.list_for_client(client_id, only_open=False)

        open_findings = [f for f in findings if f.get('status') == 'open']
        severity_counts = Counter(f.get('severity') or 'info' for f in open_findings)
        status_counts = Counter(doc.processing_status or 'sem_status' for doc in docs)
        ranking = self._build_ranking(open_findings)
        evidences = self._build_evidences(findings)
        notes_with_findings = {item['document_id'] for item in open_findings}
        attention_amount = sum(float(doc.total_amount or 0) for doc in invoices if doc.id in notes_with_findings)
        total_invoice_amount = sum(float(doc.total_amount or 0) for doc in invoices if doc.total_amount is not None)
        critical_total = severity_counts.get('critical', 0)
        warning_total = severity_counts.get('warning', 0)
        risk_score = sum(SEVERITY_WEIGHTS.get(f.get('severity'), 1) for f in open_findings)

        if critical_total:
            conclusion = 'Há apontamentos críticos. Priorize revisão técnica antes de qualquer conclusão fiscal ou contato externo.'
        elif warning_total:
            conclusion = 'Há inconsistências prováveis ou dados incompletos. Recomenda-se revisão do analista fiscal.'
        elif open_findings:
            conclusion = 'Foram encontradas observações informativas. Não há alerta crítico aberto nesta triagem.'
        else:
            conclusion = 'Nenhuma inconsistência aberta foi identificada na triagem automatizada local.'

        executive_summary = {
            'generated_at': _now_label(),
            'total_documents': len(docs),
            'total_invoices': len(invoices),
            'total_client_files': len(client_files),
            'open_findings_total': len(open_findings),
            'critical_findings': critical_total,
            'warning_findings': warning_total,
            'info_findings': severity_counts.get('info', 0),
            'documents_with_findings': len(notes_with_findings),
            'documents_without_findings': max(len(invoices) - len(notes_with_findings), 0),
            'total_invoice_amount': total_invoice_amount,
            'attention_invoice_amount': attention_amount,
            'risk_score': risk_score,
            'conclusion': conclusion,
            'method_note': (
                'Relatório gerado por motor local de triagem fiscal. Os apontamentos indicam necessidade de revisão; '
                'não representam, isoladamente, pedido de restituição, crédito definitivo ou conclusão jurídica.'
            ),
        }

        return {
            'client': _client_to_dict(client),
            'executive_summary': executive_summary,
            'severity_counts': dict(severity_counts),
            'status_counts': dict(status_counts),
            'ranking': ranking,
            'evidences': evidences,
            'findings': findings,
            'documents': [_doc_to_dict(doc) for doc in docs],
            'invoices': [_doc_to_dict(doc) for doc in invoices],
            'client_files': [_doc_to_dict(doc) for doc in client_files],
        }

    def export_client_xlsx(self, session: Session, client_id: int) -> Path:
        self.engine.auth.require(session, 'reports_export')
        data = self.build_client_report_data(session, client_id)
        client = data['client']
        output = self._output_path(client, '.xlsx')
        try:
            self._write_professional_xlsx(output, data)
        except Exception:
            output = output.with_suffix('.csv')
            self._write_fallback_csv(output, data)
        self.engine.audit.record(session.user_id, 'RELATORIO_XLSX_GERADO', 'client', client_id, str(output))
        return output

    def export_client_pdf(self, session: Session, client_id: int) -> Path:
        self.engine.auth.require(session, 'reports_export')
        data = self.build_client_report_data(session, client_id)
        client = data['client']
        output = self._output_path(client, '.pdf')
        try:
            self._write_professional_pdf(output, data)
        except Exception:
            output = output.with_suffix('.txt')
            self._write_fallback_txt(output, data)
        self.engine.audit.record(session.user_id, 'RELATORIO_PDF_GERADO', 'client', client_id, str(output))
        return output

    def export_client_csv(self, session: Session, client_id: int) -> Path:
        self.engine.auth.require(session, 'reports_export')
        data = self.build_client_report_data(session, client_id)
        client = data['client']
        output = self._output_path(client, '.csv')
        self._write_fallback_csv(output, data)
        self.engine.audit.record(session.user_id, 'RELATORIO_CSV_GERADO', 'client', client_id, str(output))
        return output

    def list_client_reports(self, session: Session, client_id: int) -> list[dict[str, Any]]:
        self.engine.auth.require(session, 'documents_read')
        prefix = f"relatorio_profissional_{client_id}_"
        reports: list[dict[str, Any]] = []
        self.engine.paths.reports.mkdir(parents=True, exist_ok=True)
        for path in sorted(self.engine.paths.reports.glob(prefix + '*'), key=lambda p: p.stat().st_mtime, reverse=True):
            reports.append({
                'name': path.name,
                'path': str(path),
                'format': path.suffix.lstrip('.').upper(),
                'size_bytes': path.stat().st_size,
                'modified_at': datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec='seconds'),
            })
        return reports

    def _output_path(self, client: dict[str, Any], suffix: str) -> Path:
        name = _safe_filename(client.get('name') or f"cliente_{client.get('id')}")
        return self.engine.paths.reports / f"relatorio_profissional_{client.get('id')}_{name}_{_timestamp()}{suffix}"

    def _build_ranking(self, findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
        buckets: dict[tuple[str, str, str], dict[str, Any]] = {}
        for item in findings:
            severity = item.get('severity') or 'info'
            code = item.get('code') or 'SEM_CODIGO'
            title = item.get('title') or code
            key = (severity, code, title)
            if key not in buckets:
                buckets[key] = {
                    'severity': severity,
                    'severity_label': SEVERITY_LABELS.get(severity, severity),
                    'code': code,
                    'title': title,
                    'count': 0,
                    'amount_total': 0.0,
                    'documents': set(),
                    'sample_message': item.get('message') or '',
                }
            bucket = buckets[key]
            bucket['count'] += 1
            bucket['amount_total'] += float(item.get('amount') or 0)
            if item.get('document_id') is not None:
                bucket['documents'].add(item.get('document_id'))
        ranking = []
        for bucket in buckets.values():
            bucket = dict(bucket)
            bucket['documents_count'] = len(bucket.pop('documents'))
            ranking.append(bucket)
        ranking.sort(key=lambda row: (SEVERITY_ORDER.get(row['severity'], 9), -row['count'], row['code']))
        return ranking

    def _build_evidences(self, findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
        ordered = sorted(
            findings,
            key=lambda item: (
                SEVERITY_ORDER.get(item.get('severity'), 9),
                _plain(item.get('original_name')),
                item.get('id') or 0,
            ),
        )
        evidence_rows: list[dict[str, Any]] = []
        for item in ordered:
            evidence_rows.append({
                'finding_id': item.get('id'),
                'document_id': item.get('document_id'),
                'severity': item.get('severity') or 'info',
                'severity_label': SEVERITY_LABELS.get(item.get('severity'), item.get('severity') or 'info'),
                'code': item.get('code') or '',
                'title': item.get('title') or '',
                'message': item.get('message') or '',
                'field': item.get('field') or '',
                'source': item.get('source') or '',
                'amount': item.get('amount'),
                'invoice_file': item.get('original_name') or '',
                'invoice_key': item.get('invoice_key') or '',
                'issue_date': item.get('issue_date') or '',
                'invoice_total': item.get('total_amount'),
                'status': item.get('status') or '',
                'review_note': item.get('review_note') or '',
            })
        return evidence_rows

    def _write_professional_xlsx(self, output: Path, data: dict[str, Any]) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        wb = Workbook()
        summary = wb.active
        summary.title = 'Resumo Executivo'
        sheets = {
            'ranking': wb.create_sheet('Ranking'),
            'evidences': wb.create_sheet('Evidências por Nota'),
            'invoices': wb.create_sheet('Notas Fiscais'),
            'documents': wb.create_sheet('Documentos'),
            'raw': wb.create_sheet('Dados Brutos'),
        }

        header_fill = PatternFill('solid', fgColor='1F4E78')
        sub_fill = PatternFill('solid', fgColor='D9EAF7')
        light_fill = PatternFill('solid', fgColor='F2F6FA')
        white_font = Font(color='FFFFFF', bold=True)
        bold_font = Font(bold=True)
        title_font = Font(size=16, bold=True)
        thin = Side(style='thin', color='D9D9D9')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        def title(ws, text: str, subtitle: str | None = None) -> None:
            ws['A1'] = text
            ws['A1'].font = title_font
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
            if subtitle:
                ws['A2'] = subtitle
                ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)

        def style_table(ws, start_row: int, end_row: int, end_col: int, name: str) -> None:
            if end_row < start_row:
                return
            for cell in ws[start_row]:
                if cell.column <= end_col:
                    cell.fill = header_fill
                    cell.font = white_font
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            for row in ws.iter_rows(min_row=start_row, max_row=end_row, max_col=end_col):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
            tab = Table(displayName=name, ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(tab)

        client = data['client']
        exec_summary = data['executive_summary']
        title(summary, 'Relatório Profissional de Análise Fiscal', 'Resumo executivo por cliente')
        summary['A4'] = 'Cliente'; summary['B4'] = client.get('name')
        summary['A5'] = 'Documento'; summary['B5'] = client.get('document')
        summary['A6'] = 'Contato'; summary['B6'] = client.get('contact_name')
        summary['A7'] = 'E-mail'; summary['B7'] = client.get('email')
        summary['A8'] = 'Gerado em'; summary['B8'] = exec_summary['generated_at']
        for cell in ['A4', 'A5', 'A6', 'A7', 'A8']:
            summary[cell].font = bold_font
            summary[cell].fill = sub_fill

        kpis = [
            ('Total de documentos', exec_summary['total_documents']),
            ('Notas fiscais', exec_summary['total_invoices']),
            ('Apontamentos abertos', exec_summary['open_findings_total']),
            ('Críticos', exec_summary['critical_findings']),
            ('Avisos', exec_summary['warning_findings']),
            ('Informativos', exec_summary['info_findings']),
            ('Notas com apontamentos', exec_summary['documents_with_findings']),
            ('Notas sem apontamentos', exec_summary['documents_without_findings']),
            ('Valor total das notas', exec_summary['total_invoice_amount']),
            ('Valor sob atenção', exec_summary['attention_invoice_amount']),
            ('Pontuação de risco local', exec_summary['risk_score']),
        ]
        summary.append([])
        summary.append(['Indicador', 'Valor'])
        start = summary.max_row
        for label, value in kpis:
            summary.append([label, value])
        style_table(summary, start, summary.max_row, 2, 'TabelaResumoExecutivo')
        for row in range(start + 1, summary.max_row + 1):
            if 'Valor' in str(summary.cell(row=row, column=1).value):
                summary.cell(row=row, column=2).number_format = 'R$ #,##0.00'

        note_row = summary.max_row + 2
        summary.cell(note_row, 1, 'Conclusão executiva').font = bold_font
        summary.cell(note_row, 1).fill = sub_fill
        summary.cell(note_row, 2, exec_summary['conclusion'])
        summary.cell(note_row + 1, 1, 'Nota metodológica').font = bold_font
        summary.cell(note_row + 1, 1).fill = sub_fill
        summary.cell(note_row + 1, 2, exec_summary['method_note'])
        summary.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=8)
        summary.merge_cells(start_row=note_row + 1, start_column=2, end_row=note_row + 1, end_column=8)
        for row in range(note_row, note_row + 2):
            for col in range(1, 9):
                summary.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical='top')
                summary.cell(row=row, column=col).border = border

        ranking_ws = sheets['ranking']
        title(ranking_ws, 'Ranking de Inconsistências', 'Agrupamento por gravidade, código e título')
        ranking_headers = ['Gravidade', 'Código', 'Título', 'Ocorrências', 'Notas afetadas', 'Valor relacionado', 'Exemplo de mensagem']
        ranking_ws.append([])
        ranking_ws.append(ranking_headers)
        ranking_start = ranking_ws.max_row
        for row in data['ranking']:
            ranking_ws.append([
                row['severity_label'], row['code'], row['title'], row['count'], row['documents_count'],
                row['amount_total'], row['sample_message'],
            ])
        if ranking_ws.max_row == ranking_start:
            ranking_ws.append(['-', '-', 'Nenhuma inconsistência aberta', 0, 0, 0, ''])
        style_table(ranking_ws, ranking_start, ranking_ws.max_row, len(ranking_headers), 'TabelaRankingInconsistencias')
        for row in range(ranking_start + 1, ranking_ws.max_row + 1):
            ranking_ws.cell(row=row, column=6).number_format = 'R$ #,##0.00'

        evidence_ws = sheets['evidences']
        title(evidence_ws, 'Evidências por Nota', 'Detalhamento para conferência do analista')
        evidence_headers = ['Gravidade', 'Nota/Arquivo', 'Chave fiscal', 'Emissão', 'Valor da nota', 'Código', 'Título', 'Mensagem', 'Campo', 'Fonte', 'Status revisão', 'Observação', 'Valor relacionado']
        evidence_ws.append([])
        evidence_ws.append(evidence_headers)
        evidence_start = evidence_ws.max_row
        for row in data['evidences']:
            evidence_ws.append([
                row['severity_label'], row['invoice_file'], row['invoice_key'], row['issue_date'], row['invoice_total'],
                row['code'], row['title'], row['message'], row['field'], row['source'], row.get('status'), row.get('review_note'), row['amount'],
            ])
        if evidence_ws.max_row == evidence_start:
            evidence_ws.append(['-', '-', '-', '-', 0, '-', 'Nenhuma inconsistência aberta', '', '', '', 0])
        style_table(evidence_ws, evidence_start, evidence_ws.max_row, len(evidence_headers), 'TabelaEvidenciasNota')
        for row in range(evidence_start + 1, evidence_ws.max_row + 1):
            evidence_ws.cell(row=row, column=5).number_format = 'R$ #,##0.00'
            evidence_ws.cell(row=row, column=13).number_format = 'R$ #,##0.00'

        invoices_ws = sheets['invoices']
        title(invoices_ws, 'Notas Fiscais', 'Documentos fiscais importados para o cliente')
        invoice_headers = ['ID', 'Arquivo', 'Tipo', 'Chave fiscal', 'Emissão', 'Emitente', 'Documento emitente', 'Destinatário', 'Valor', 'Leitura', 'Situação']
        invoices_ws.append([])
        invoices_ws.append(invoice_headers)
        invoice_start = invoices_ws.max_row
        for doc in data['invoices']:
            invoices_ws.append([
                doc.get('id'), doc.get('original_name'), doc.get('file_type'), doc.get('invoice_key'), doc.get('issue_date'),
                doc.get('issuer_name'), doc.get('issuer_document'), doc.get('recipient_name'), doc.get('total_amount'),
                doc.get('parser_status'), doc.get('processing_status'),
            ])
        if invoices_ws.max_row == invoice_start:
            invoices_ws.append(['', 'Nenhuma nota fiscal importada', '', '', '', '', '', '', 0, '', ''])
        style_table(invoices_ws, invoice_start, invoices_ws.max_row, len(invoice_headers), 'TabelaNotasFiscais')
        for row in range(invoice_start + 1, invoices_ws.max_row + 1):
            invoices_ws.cell(row=row, column=9).number_format = 'R$ #,##0.00'

        docs_ws = sheets['documents']
        title(docs_ws, 'Documentos do Cliente', 'Contratos e documentos não fiscais')
        doc_headers = ['ID', 'Categoria', 'Arquivo', 'Tipo', 'Situação', 'Caminho interno']
        docs_ws.append([])
        docs_ws.append(doc_headers)
        docs_start = docs_ws.max_row
        for doc in data['client_files']:
            docs_ws.append([doc.get('id'), doc.get('category'), doc.get('original_name'), doc.get('file_type'), doc.get('processing_status'), doc.get('storage_path')])
        if docs_ws.max_row == docs_start:
            docs_ws.append(['', 'Nenhum documento complementar importado', '', '', '', ''])
        style_table(docs_ws, docs_start, docs_ws.max_row, len(doc_headers), 'TabelaDocumentosCliente')

        raw_ws = sheets['raw']
        title(raw_ws, 'Dados Brutos dos Apontamentos', 'Base completa exportada para filtros e auditoria interna')
        raw_headers = ['ID', 'Documento ID', 'Gravidade', 'Código', 'Título', 'Mensagem', 'Valor', 'Campo', 'Fonte', 'Status', 'Observação revisão', 'Arquivo', 'Chave fiscal', 'Data de emissão']
        raw_ws.append([])
        raw_ws.append(raw_headers)
        raw_start = raw_ws.max_row
        for item in data['findings']:
            raw_ws.append([
                item.get('id'), item.get('document_id'), item.get('severity'), item.get('code'), item.get('title'),
                item.get('message'), item.get('amount'), item.get('field'), item.get('source'), item.get('status'), item.get('review_note'),
                item.get('original_name'), item.get('invoice_key'), item.get('issue_date'),
            ])
        if raw_ws.max_row == raw_start:
            raw_ws.append(['', '', '', '', 'Nenhum apontamento registrado', '', 0, '', '', '', '', '', '', ''])
        style_table(raw_ws, raw_start, raw_ws.max_row, len(raw_headers), 'TabelaDadosBrutos')
        for row in range(raw_start + 1, raw_ws.max_row + 1):
            raw_ws.cell(row=row, column=7).number_format = 'R$ #,##0.00'

        widths = {
            'Resumo Executivo': {'A': 26, 'B': 22, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 16, 'H': 16},
            'Ranking': {'A': 14, 'B': 26, 'C': 38, 'D': 14, 'E': 14, 'F': 16, 'G': 70},
            'Evidências por Nota': {'A': 14, 'B': 28, 'C': 48, 'D': 20, 'E': 15, 'F': 26, 'G': 32, 'H': 80, 'I': 20, 'J': 12, 'K': 16, 'L': 32, 'M': 16},
            'Notas Fiscais': {'A': 8, 'B': 30, 'C': 10, 'D': 48, 'E': 20, 'F': 35, 'G': 20, 'H': 35, 'I': 16, 'J': 14, 'K': 14},
            'Documentos': {'A': 8, 'B': 18, 'C': 34, 'D': 10, 'E': 16, 'F': 70},
            'Dados Brutos': {'A': 8, 'B': 12, 'C': 12, 'D': 26, 'E': 32, 'F': 80, 'G': 14, 'H': 20, 'I': 12, 'J': 12, 'K': 32, 'L': 30, 'M': 48, 'N': 20},
        }
        for ws in wb.worksheets:
            ws.freeze_panes = 'A4'
            ws.sheet_view.showGridLines = False
            for col_letter, width in widths.get(ws.title, {}).items():
                ws.column_dimensions[col_letter].width = width
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cell.alignment = Alignment(vertical='top', wrap_text=True)
            for row_idx in range(1, ws.max_row + 1):
                ws.row_dimensions[row_idx].height = 32 if row_idx > 2 else 24
            # destaque visual leve para linhas de dados críticos/atenção
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
                first = str(row[0].value or '').lower()
                if 'crítico' in first or first == 'critical':
                    for cell in row:
                        cell.fill = PatternFill('solid', fgColor='FCE4D6')
                elif 'atenção' in first or first == 'warning':
                    for cell in row:
                        if cell.fill.fill_type is None:
                            cell.fill = light_fill

        wb.save(output)

    def _write_professional_pdf(self, output: Path, data: dict[str, Any]) -> None:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('RecuperaTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=18, leading=22, alignment=TA_CENTER, spaceAfter=12)
        subtitle_style = ParagraphStyle('RecuperaSubtitle', parent=styles['Normal'], fontSize=9, leading=12, alignment=TA_CENTER, textColor=colors.HexColor('#444444'), spaceAfter=18)
        h2 = ParagraphStyle('RecuperaH2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=12, leading=15, spaceBefore=8, spaceAfter=6, textColor=colors.HexColor('#1F4E78'))
        body = ParagraphStyle('RecuperaBody', parent=styles['BodyText'], fontSize=8, leading=10, alignment=TA_LEFT)
        small = ParagraphStyle('RecuperaSmall', parent=styles['BodyText'], fontSize=7, leading=9, alignment=TA_LEFT)

        doc = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=1.4 * cm, leftMargin=1.4 * cm, topMargin=1.4 * cm, bottomMargin=1.4 * cm)
        story: list[Any] = []
        client = data['client']
        exec_summary = data['executive_summary']

        story.append(Paragraph('Relatório Profissional de Análise Fiscal', title_style))
        story.append(Paragraph('Resumo executivo, ranking de inconsistências e evidências por nota', subtitle_style))
        story.append(self._pdf_table([
            ['Cliente', _plain(client.get('name'))],
            ['Documento', _plain(client.get('document'))],
            ['Contato', _plain(client.get('contact_name'))],
            ['E-mail', _plain(client.get('email'))],
            ['Gerado em', exec_summary['generated_at']],
        ], col_widths=[4 * cm, 12 * cm]))
        story.append(Spacer(1, 0.35 * cm))
        story.append(Paragraph('Resumo executivo', h2))
        story.append(self._pdf_table([
            ['Indicador', 'Valor'],
            ['Documentos totais', exec_summary['total_documents']],
            ['Notas fiscais importadas', exec_summary['total_invoices']],
            ['Apontamentos abertos', exec_summary['open_findings_total']],
            ['Críticos', exec_summary['critical_findings']],
            ['Avisos', exec_summary['warning_findings']],
            ['Informativos', exec_summary['info_findings']],
            ['Notas com apontamentos', exec_summary['documents_with_findings']],
            ['Valor total das notas', _money(exec_summary['total_invoice_amount'])],
            ['Valor sob atenção', _money(exec_summary['attention_invoice_amount'])],
            ['Pontuação de risco local', exec_summary['risk_score']],
        ], header=True, col_widths=[8 * cm, 8 * cm]))
        story.append(Spacer(1, 0.2 * cm))
        story.append(Paragraph(f"<b>Conclusão:</b> {exec_summary['conclusion']}", body))
        story.append(Spacer(1, 0.15 * cm))
        story.append(Paragraph(f"<b>Nota metodológica:</b> {exec_summary['method_note']}", small))

        story.append(PageBreak())
        story.append(Paragraph('Ranking de inconsistências', h2))
        ranking_rows = [['Gravidade', 'Código', 'Ocorrências', 'Notas', 'Título']]
        for row in data['ranking'][:25]:
            ranking_rows.append([
                row['severity_label'], row['code'], row['count'], row['documents_count'],
                Paragraph(_plain(row['title']), small),
            ])
        if len(ranking_rows) == 1:
            ranking_rows.append(['-', '-', 0, 0, 'Nenhuma inconsistência aberta'])
        story.append(self._pdf_table(ranking_rows, header=True, col_widths=[2.2 * cm, 3.6 * cm, 2.0 * cm, 1.5 * cm, 6.5 * cm]))

        story.append(Spacer(1, 0.3 * cm))
        story.append(Paragraph('Evidências por nota', h2))
        evidence_rows = [['Gravidade', 'Arquivo', 'Código', 'Mensagem']]
        for row in data['evidences'][:80]:
            evidence_rows.append([
                row['severity_label'],
                Paragraph(_plain(row['invoice_file'])[:80], small),
                row['code'],
                Paragraph(_plain(row['message']), small),
            ])
        if len(evidence_rows) == 1:
            evidence_rows.append(['-', '-', '-', 'Nenhuma inconsistência aberta'])
        story.append(self._pdf_table(evidence_rows, header=True, col_widths=[2.1 * cm, 4.1 * cm, 3.3 * cm, 6.5 * cm]))

        if len(data['evidences']) > 80:
            story.append(Spacer(1, 0.2 * cm))
            story.append(Paragraph(f"O PDF mostra as primeiras 80 evidências. A exportação Excel contém todas as {len(data['evidences'])} evidências.", small))

        story.append(PageBreak())
        story.append(Paragraph('Notas fiscais importadas', h2))
        invoice_rows = [['Arquivo', 'Chave fiscal', 'Emissão', 'Valor', 'Situação']]
        for doc_item in data['invoices'][:80]:
            invoice_rows.append([
                Paragraph(_plain(doc_item.get('original_name'))[:70], small),
                Paragraph(_plain(doc_item.get('invoice_key'))[:55], small),
                _plain(doc_item.get('issue_date'))[:19],
                _money(doc_item.get('total_amount')),
                _plain(doc_item.get('processing_status')),
            ])
        if len(invoice_rows) == 1:
            invoice_rows.append(['Nenhuma nota fiscal importada', '', '', '', ''])
        story.append(self._pdf_table(invoice_rows, header=True, col_widths=[4.0 * cm, 5.0 * cm, 2.5 * cm, 2.2 * cm, 2.3 * cm]))
        story.append(Spacer(1, 0.25 * cm))
        story.append(Paragraph('Este relatório é material de apoio para conferência interna do analista. A conclusão fiscal definitiva depende de validação documental, legislação aplicável, enquadramento do cliente e revisão técnica.', small))

        doc.build(story)

    def _pdf_table(self, rows: list[list[Any]], header: bool = False, col_widths: list[float] | None = None):
        from reportlab.lib import colors
        from reportlab.platypus import Table, TableStyle

        table = Table(rows, colWidths=col_widths, repeatRows=1 if header else 0)
        style = [
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#D9D9D9')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]
        if header:
            style.extend([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ])
        else:
            style.extend([
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#D9EAF7')),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ])
        table.setStyle(TableStyle(style))
        return table

    def _write_fallback_csv(self, output: Path, data: dict[str, Any]) -> None:
        with output.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f, delimiter=';')
            writer.writerow(['gravidade', 'arquivo', 'chave_fiscal', 'codigo', 'titulo', 'mensagem', 'valor'])
            for row in data['evidences']:
                writer.writerow([
                    row['severity_label'], row['invoice_file'], row['invoice_key'], row['code'], row['title'], row['message'], row['amount'],
                ])

    def _write_fallback_txt(self, output: Path, data: dict[str, Any]) -> None:
        lines = [
            'Relatório Profissional de Análise Fiscal',
            f"Cliente: {data['client'].get('name')}",
            f"Gerado em: {data['executive_summary']['generated_at']}",
            '',
            'Resumo executivo:',
            data['executive_summary']['conclusion'],
            '',
            'Evidências:',
        ]
        for row in data['evidences']:
            lines.append(f"- {row['severity_label']} | {row['invoice_file']} | {row['code']} | {row['message']}")
        if not data['evidences']:
            lines.append('- Nenhuma inconsistência aberta.')
        output.write_text('\n'.join(lines), encoding='utf-8')
